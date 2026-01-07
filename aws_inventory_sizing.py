import boto3
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
# ---------------- Configuration ----------------

REGION = "ap-southeast-1"

def safe(fn, default=None):
    try:
        return fn()
    except Exception:
        return default

# ---------------- Excel Styling ----------------
def style_sheet(ws):
    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    # Header row
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.border = thin_border
        ws.column_dimensions[cell.column_letter].width = 25
    # Border all cells
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border

# ---------------- EC2 ----------------
def fetch_ec2_instance_type_info():
    client = boto3.client("ec2", region_name=REGION)
    info = {}
    paginator = client.get_paginator("describe_instance_types")
    for page in paginator.paginate():
        for inst in page["InstanceTypes"]:
            info[inst["InstanceType"]] = {
                "vCPU": inst["VCpuInfo"]["DefaultVCpus"],
                "MemoryMB": inst["MemoryInfo"]["SizeInMiB"],
                "NetworkPerformance": inst.get("NetworkInfo", {}).get("NetworkPerformance"),
                "EBS_Optimized": inst.get("EbsInfo", {}).get("EbsOptimizedSupport"),
                "MaxBandwidth": inst.get("NetworkInfo", {}).get("MaximumNetworkBandwidthInGbps"),
            }
    return info

def fetch_ec2():
    ec2 = boto3.client("ec2", region_name=REGION)
    typeinfo = fetch_ec2_instance_type_info()
    rows = []
    paginator = ec2.get_paginator("describe_instances")
    for page in paginator.paginate():
        for res in page["Reservations"]:
            for inst in res.get("Instances", []):
                state = inst["State"]["Name"]
                if state not in ["running", "stopped"]:
                    continue
                inst_type = inst.get("InstanceType")
                details = typeinfo.get(inst_type, {})
                sg = ",".join([s["GroupName"] for s in inst.get("SecurityGroups", [])])
                rows.append({
                    "InstanceId": inst.get("InstanceId"),
                    "Name": next((t["Value"] for t in inst.get("Tags", []) if t["Key"] == "Name"), ""),
                    "InstanceType": inst_type,
                    "vCPU": details.get("vCPU"),
                    "MemoryMB": details.get("MemoryMB"),
                    "NetworkPerformance": details.get("NetworkPerformance"),
                    "MaxBandwidthGbps": details.get("MaxBandwidth"),
                    "EBSOptimized": details.get("EBS_Optimized"),
                    "AMI": inst.get("ImageId"),
                    "Architecture": inst.get("Architecture"),
                    "Hypervisor": inst.get("Hypervisor"),
                    "KeyName": inst.get("KeyName"),
                    "SubnetId": inst.get("SubnetId"),
                    "VPCId": inst.get("VpcId"),
                    "PrivateIP": inst.get("PrivateIpAddress"),
                    "PublicIP": inst.get("PublicIpAddress"),
                    "RootDeviceType": inst.get("RootDeviceType"),
                    "RootDeviceName": inst.get("RootDeviceName"),
                    "Platform": inst.get("PlatformDetails"),
                    "Tenancy": inst.get("Placement", {}).get("Tenancy"),
                    "State": state,
                    "LaunchTime": str(inst.get("LaunchTime")),
                    "SecurityGroups": sg,
                })
    return rows

# ---------------- EKS ----------------
def fetch_eks_full():
    client = boto3.client("eks", region_name=REGION)
    rows = []
    clusters = safe(lambda: client.list_clusters()["clusters"], [])
    for c in clusters:
        cdesc = client.describe_cluster(name=c)["cluster"]
        base = {
            "ClusterName": c,
            "ClusterVersion": cdesc.get("version"),
            "ClusterStatus": cdesc.get("status"),
            "Endpoint": cdesc.get("endpoint"),
            "ClusterRoleArn": cdesc.get("roleArn"),
            "VpcId": cdesc.get("resourcesVpcConfig", {}).get("vpcId"),
        }
        ngs = client.list_nodegroups(clusterName=c)["nodegroups"]
        for ng in ngs:
            ngd = client.describe_nodegroup(clusterName=c, nodegroupName=ng)["nodegroup"]
            rows.append({
                **base,
                "NodeGroup": ng,
                "InstanceTypes": ",".join(ngd.get("instanceTypes", [])),
                "AMIType": ngd.get("amiType"),
                "CapacityType": ngd.get("capacityType"),
                "DesiredSize": ngd["scalingConfig"]["desiredSize"],
                "MinSize": ngd["scalingConfig"]["minSize"],
                "MaxSize": ngd["scalingConfig"]["maxSize"],
                "NodeRole": ngd.get("nodeRole"),
                "DiskSizeGB": ngd.get("diskSize"),
            })
    return rows

# ---------------- EFS ----------------
def fetch_efs():
    efs = boto3.client("efs", region_name=REGION)
    cw = boto3.client("cloudwatch", region_name=REGION)
    rows = []
    for fs in efs.describe_file_systems()["FileSystems"]:
        size = 0
        data = cw.get_metric_statistics(
            Namespace="AWS/EFS",
            MetricName="FileSystemSize",
            Dimensions=[{"Name": "FileSystemId", "Value": fs["FileSystemId"]}],
            StartTime=datetime.utcnow() - timedelta(hours=2),
            EndTime=datetime.utcnow(),
            Period=3600,
            Statistics=["Average"]
        )
        if data["Datapoints"]:
            size = data["Datapoints"][0]["Average"]
        rows.append({
            "FileSystemId": fs["FileSystemId"],
            "PerformanceMode": fs["PerformanceMode"],
            "Encrypted": fs["Encrypted"],
            "ThroughputMode": fs.get("ThroughputMode"),
            "LifecyclePolicies": str(fs.get("LifecyclePolicies")),
            "SizeGB": round(size / (1024 ** 3), 2),
        })
    return rows

# ---------------- MSK ----------------
def fetch_msk():
    client = boto3.client("kafka", region_name=REGION)
    rows = []
    clusters = safe(lambda: client.list_clusters()["ClusterInfoList"], [])
    for c in clusters:
        cid = c["ClusterArn"]
        desc = client.describe_cluster(ClusterArn=cid)["ClusterInfo"]
        rows.append({
            "ClusterName": desc["ClusterName"],
            "ClusterArn": desc.get("ClusterArn"),
            "State": desc.get("State"),
            "CreationTime": str(desc.get("CreationTime")),
            "NumberOfBrokers": desc.get("NumberOfBrokerNodes"),
            "InstanceType": desc.get("BrokerNodeGroupInfo", {}).get("InstanceType"),
            "StoragePerBrokerGB": desc.get("BrokerNodeGroupInfo", {}).get("StorageInfo", {}).get("EbsStorageInfo", {}).get("VolumeSize"),
            "TotalStorageGB": desc.get("NumberOfBrokerNodes") * desc.get("BrokerNodeGroupInfo", {}).get("StorageInfo", {}).get("EbsStorageInfo", {}).get("VolumeSize",0),
            "ClientAuthentication": str(desc.get("ClientAuthentication")),
            "EncryptionInTransit": str(desc.get("EncryptionInfo", {}).get("EncryptionInTransit")),
            "EncryptionAtRestKMSKey": desc.get("EncryptionInfo", {}).get("EncryptionAtRest", {}).get("DataVolumeKMSKeyId"),
            "ZookeeperConnectString": desc.get("ZookeeperConnectString"),
            "ZookeeperConnectStringTLS": desc.get("ZookeeperConnectStringTLS"),
            "BrokerAZDistribution": str(desc.get("BrokerNodeGroupInfo", {}).get("BrokerAZDistribution")),
            "EnhancedMonitoring": desc.get("EnhancedMonitoring"),
            "OpenMonitoring": str(desc.get("OpenMonitoring")),
            "LoggingInfo": str(desc.get("LoggingInfo")),
            "PublicAccess": str(desc.get("PublicAccess", {})),
            "CurrentKafkaVersion": desc.get("CurrentBrokerSoftwareInfo", {}).get("KafkaVersion"),
            "TargetKafkaVersion": desc.get("TargetKafkaVersion"),
            "Tags": str(desc.get("Tags", {}))
        })
    return rows

# ---------------- RDS ----------------
def fetch_rds():
    client = boto3.client("rds", region_name=REGION)
    rows = []
    for db in client.describe_db_instances()["DBInstances"]:
        if db["DBInstanceStatus"] not in ["available"]:
            continue
        rows.append({
            "Identifier": db["DBInstanceIdentifier"],
            "ARN": db.get("DBInstanceArn"),
            "Engine": db.get("Engine"),
            "EngineVersion": db.get("EngineVersion"),
            "DBClusterIdentifier": db.get("DBClusterIdentifier"),
            "Class": db.get("DBInstanceClass"),
            "StorageGB": db.get("AllocatedStorage"),
            "StorageType": db.get("StorageType"),
            "IOPS": db.get("Iops"),
            "MaxAllocatedStorage": db.get("MaxAllocatedStorage"),
            "MultiAZ": db.get("MultiAZ"),
            "AvailabilityZone": db.get("AvailabilityZone"),
            "SecondaryAZs": ",".join(db.get("SecondaryAvailabilityZones", [])),
            "BackupRetentionDays": db.get("BackupRetentionPeriod"),
            "AutoMinorVersionUpgrade": db.get("AutoMinorVersionUpgrade"),
            "PubliclyAccessible": db.get("PubliclyAccessible"),
            "PreferredMaintenanceWindow": db.get("PreferredMaintenanceWindow"),
            "PreferredBackupWindow": db.get("PreferredBackupWindow"),
            "DeletionProtection": db.get("DeletionProtection"),
            "MonitoringInterval": db.get("MonitoringInterval"),
            "EnhancedMonitoringRoleArn": db.get("MonitoringRoleArn"),
            "PerformanceInsightsEnabled": db.get("PerformanceInsightsEnabled"),
            "PerformanceInsightsRetentionPeriod": db.get("PerformanceInsightsRetentionPeriod"),
            "PerformanceInsightsKMSKeyId": db.get("PerformanceInsightsKMSKeyId"),
            "KmsKeyId": db.get("KmsKeyId"),
            "StorageEncrypted": db.get("StorageEncrypted"),
            "ReadReplicaCount": len(db.get("ReadReplicaDBInstanceIdentifiers", [])),
            "Endpoint": db.get("Endpoint", {}).get("Address"),
            "DBSubnetGroup": db.get("DBSubnetGroup", {}).get("DBSubnetGroupName"),
            "LicenseModel": db.get("LicenseModel"),
            "Tags": str(safe(lambda: client.list_tags_for_resource(ResourceName=db.get("DBInstanceArn"))["TagList"], {}))
        })
    return rows

# ---------------- Excel Writer ----------------
def write_sheet(wb, name, rows):
    if not rows:
        return
    ws = wb.create_sheet(name)
    headers = list(rows[0].keys())
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    style_sheet(ws)

# ---------------- MAIN ----------------
def main():
    # Get account ID
    sts = boto3.client("sts")
    account_id = sts.get_caller_identity()["Account"]

    timestamp = datetime.now().strftime("%d%m%y%H%M%S")
    filename = f"aws_inventory_{account_id}_{timestamp}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)

    write_sheet(wb, "EC2", fetch_ec2())
    write_sheet(wb, "EKS", fetch_eks_full())
    write_sheet(wb, "EFS", fetch_efs())
    write_sheet(wb, "MSK", fetch_msk())
    write_sheet(wb, "RDS", fetch_rds())

    wb.save(filename)
    print("Inventory generated:", filename)

if __name__ == "__main__":
    main()
